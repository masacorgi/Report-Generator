import os
import sys
from tkinter import Tk, Label, Button, Entry, filedialog, messagebox
from tkinter import ttk  # ✅ 추가
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

# =========================
# 실행 경로
# =========================
def get_base_path():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

folder_path = ""

# =========================
# 폰트
# =========================
def set_korean_font(run, font_name="바탕"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

# =========================
# 표 테두리
# =========================
def set_table_border(table):
    tbl = table._element
    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '2')
        border.set(qn('w:color'), '000000')
        borders.append(border)

    tblPr.append(borders)

# =========================
# 이미지 크기 계산
# =========================
def calc_size(img_path, max_w, max_h):
    img = Image.open(img_path)
    w, h = img.size
    ratio = min(max_w / w, max_h / h)
    return w * ratio, h * ratio

# =========================
# 이미지 삽입
# =========================
def insert_image(table, row, img_path, cell_height):
    MAX_W = 6.3
    MAX_H = cell_height

    w, h = calc_size(img_path, MAX_W, MAX_H)

    cell = table.cell(row, 0)
    cell.height = Inches(cell_height)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run()
    run.add_picture(img_path, width=Inches(w), height=Inches(h))

    # 캡션
    caption = table.cell(row+1, 0)
    caption.height = Inches(0.5)

    p2 = caption.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ✅ 확장자 제거
    filename_only = os.path.splitext(os.path.basename(img_path))[0]

    run2 = p2.add_run(filename_only)
    run2.font.size = Pt(15)
    set_korean_font(run2)

# =========================
# 문서 생성
# =========================
def generate_doc():
    if not folder_path:
        messagebox.showerror("오류", "사진 폴더 선택 필요")
        return

    filename = entry_name.get().strip()
    if not filename:
        messagebox.showerror("오류", "파일명을 입력하세요")
        return

    save_path = os.path.join(get_base_path(), filename + ".docx")

    images = [f for f in os.listdir(folder_path)
              if f.lower().endswith(('.jpg', '.png', '.jpeg'))]

    if not images:
        messagebox.showerror("오류", "이미지 없음")
        return

    images.sort()

    # ✅ 진행바 설정
    progress['maximum'] = len(images)
    progress['value'] = 0

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # 제목
    title = doc.add_paragraph()
    run = title.add_run("모바일알람 설치 사진")
    run.bold = True
    run.font.size = Pt(20)
    set_korean_font(run)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    IMAGE_HEIGHT = 4.0

    count = 0  # 진행률 카운트

    # 첫 페이지
    table = doc.add_table(rows=4, cols=1)
    set_table_border(table)

    insert_image(table, 0, os.path.join(folder_path, images[0]), IMAGE_HEIGHT)
    count += 1
    progress['value'] = count
    root.update_idletasks()

    if len(images) > 1:
        insert_image(table, 2, os.path.join(folder_path, images[1]), IMAGE_HEIGHT)
        count += 1
        progress['value'] = count
        root.update_idletasks()

    if len(images) > 2:
        doc.add_page_break()

    # 이후 페이지
    for i in range(2, len(images), 2):
        table = doc.add_table(rows=4, cols=1)
        set_table_border(table)

        insert_image(table, 0, os.path.join(folder_path, images[i]), IMAGE_HEIGHT)
        count += 1
        progress['value'] = count
        root.update_idletasks()

        if i + 1 < len(images):
            insert_image(table, 2, os.path.join(folder_path, images[i+1]), IMAGE_HEIGHT)
            count += 1
            progress['value'] = count
            root.update_idletasks()

        if i + 2 < len(images):
            doc.add_page_break()

    doc.save(save_path)

    progress['value'] = 0  # 초기화
    messagebox.showinfo("완료", f"저장 완료:\n{save_path}")

# =========================
# UI
# =========================
def select_folder():
    global folder_path
    folder_path = filedialog.askdirectory()
    label_folder.config(text=folder_path)

root = Tk()
root.title("설치사진 보고서 생성기")

# 창 크기
root.geometry("500x400")

# 1️⃣ 상단 여백
Label(root, text="").pack(pady=5)

# 2️⃣ 사진 폴더 선택
Label(root, text="사진 폴더 선택").pack()
Button(root, text="폴더 선택", command=select_folder).pack()

label_folder = Label(root, text="")
label_folder.pack()

# 3️⃣ 여백
Label(root, text="").pack(pady=5)

# 4️⃣ 파일명 입력
Label(root, text="파일명 입력").pack()
entry_name = Entry(root, width=30)
entry_name.pack()

# 5️⃣ 여백
Label(root, text="").pack(pady=5)

# 6️⃣ 안내문구 (2줄)
Label(
    root,
    text="* 이미지는 파일명 이름 순으로 삽입됩니다. *\n파일명 앞에 숫자를 붙여 조정해주세요.",
    justify="center"
).pack()

Label(
    root,
    text="* 보고서 파일은 보고서 생성기와 동일한 폴더에 생성됩니다. *",
    justify="center"
).pack(pady=(5, 10))

Label(root, text="").pack(pady=5)

# 7️⃣ 생성 버튼
Button(root, text="생성", command=generate_doc).pack()

# 8️⃣ 진행바
progress = ttk.Progressbar(root, orient="horizontal", length=300, mode="determinate")
progress.pack(pady=10)

root.mainloop()
